#!/usr/bin/env python3
"""
Parse the ministry Excel catalogue into src/data JSON.

Same rules as scripts/parse-excel.ps1:
  - columns by header name (extra MI columns do not shift Kategorie)
  - type from the correct answer: T/N = basic, A/B/C = specialist
    (Excel "Zakres struktury" is sometimes wrong)
  - keep media file names even if the local pack is missing (CDN)
  - --drop-missing-media only when asked: then missing raw files clear media
  - fill missing EN/DE/UK via translate-questions.py unless --skip-translate-gaps
    (Gemini from --gemini-api-key, else .geminienv / GEMINI_API_KEY, else Google)
  - keep an existing translation when the new Excel has no text for that language
  - ministry text replaces it
  - if the Polish question or A/B/C changed (case, spacing, and dash/quote
    shape ignored), drop the old translation so the gap is filled again
  - category files list question ids only; facts live once in questions.json;
    question and answer text, including Polish, lives in translations_*.json

Usage:
  python3 scripts/parse-excel.py
  python3 scripts/parse-excel.py --out-dir src/data
"""

import argparse
import importlib.util
import json
import os
import re
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is required: pip install openpyxl")

CATEGORIES = ["A", "A1", "A2", "AM", "B", "B1", "C", "C1", "D", "D1", "PT", "T"]

EXAM_RULES = {
    "totalQuestions": 32,
    "basicQuestions": 20,
    "specialistQuestions": 12,
    "maxPoints": 74,
    "passThreshold": 68,
    "totalTimeSeconds": 1500,
    "basicTimeSeconds": 20,
    "specialistTimeSeconds": 50,
    "basicPoints": [3, 3, 3, 3, 3, 3, 3, 3, 3, 3, 2, 2, 2, 2, 2, 2, 1, 1, 1, 1],
    "specialistPoints": [3, 3, 3, 3, 3, 3, 2, 2, 2, 2, 1, 1],
}

MEDIA_EXT_MAP = {
    ".wmv": (".mp4", "video"),
    ".jpg": (".webp", "image"),
    ".jpeg": (".webp", "image"),
}


def cell(row, idx):
    if idx is None or idx < 0 or idx >= len(row) or row[idx] is None:
        return ""
    return str(row[idx]).strip()


def find_header(header, label, exact=None, regex=None, required=True):
    for i, raw in enumerate(header):
        h = "" if raw is None else str(raw).strip()
        if exact:
            for name in exact:
                if name and h == name:
                    return i
        if regex and re.search(regex, h):
            return i
    if required:
        names = " | ".join("" if x is None else str(x) for x in header)
        sys.exit(f"Missing column {label}. Headers: {names}")
    return -1


def question_type(structure, correct, ans_a, ans_b, ans_c):
    if correct in ("A", "B", "C"):
        return "specialist"
    if correct in ("T", "N"):
        return "basic"
    if ans_a and ans_b and ans_c:
        return "specialist"
    if structure == "PODSTAWOWY":
        return "basic"
    return "specialist"


def media_lookup(media_dir: Path | None):
    if media_dir is None or not media_dir.exists():
        return None
    names = {}
    for f in media_dir.iterdir():
        if f.is_file():
            names[f.name.lower()] = True
    return names


def resolve_media(raw_filename, lookup, drop_missing):
    if not raw_filename:
        return None, None, False
    src_ext = os.path.splitext(raw_filename)[1].lower()
    mapping = MEDIA_EXT_MAP.get(src_ext)
    if mapping is None:
        print(f"  WARNING: unknown media extension '{src_ext}' for '{raw_filename}'")
        media_name, media_type = raw_filename, "unknown"
    else:
        target_ext, media_type = mapping
        media_name = os.path.splitext(raw_filename)[0] + target_ext
    if drop_missing and lookup is not None and raw_filename.lower() not in lookup:
        return None, None, True
    return media_name, media_type, False


def normalize_polish_text(text) -> str:
    if not text:
        return ""
    s = str(text).replace("\u00a0", " ").replace("\u202f", " ").replace("\u2009", " ")
    s = s.translate(str.maketrans({
        "\u2010": "-", "\u2011": "-", "\u2012": "-", "\u2013": "-",
        "\u2014": "-", "\u2015": "-", "\u2212": "-",
        "\u201e": '"', "\u201c": '"', "\u201d": '"', "\u00ab": '"', "\u00bb": '"',
        "\u2018": "'", "\u2019": "'",
    }))
    s = s.lower()
    return re.sub(r"\s+", " ", s).strip()


def polish_question_key(question) -> str:
    if isinstance(question, dict):
        get = question.get
    else:
        get = lambda name: getattr(question, name, None)
    parts = [normalize_polish_text(get("q"))]
    for field in ("a", "b", "c"):
        val = get(field)
        if val:
            parts.append(normalize_polish_text(val))
    return "\n".join(parts)


def changed_polish_ids(old_keys: dict, questions) -> set[str]:
    changed = set()
    seen = set()
    for question in questions:
        qid = str(question.get("id") if isinstance(question, dict) else question.id)
        if not qid or qid in seen:
            continue
        seen.add(qid)
        if qid in old_keys and old_keys[qid] != polish_question_key(question):
            changed.add(qid)
    return changed


def question_fact(question: dict) -> dict:
    fact = {
        "type": question.get("type"),
        "correct": question.get("correct"),
        "points": question.get("points"),
    }
    if question.get("media"):
        fact["media"] = question.get("media")
        fact["mediaType"] = question.get("mediaType")
    return fact


def polish_record(question: dict) -> dict:
    text = {"q": question.get("q") or "", "src": "excel"}
    if question.get("a") or question.get("b") or question.get("c"):
        text["a"] = question.get("a") or ""
        text["b"] = question.get("b") or ""
        text["c"] = question.get("c") or ""
    return text


def load_existing_polish_keys(out_dir: Path) -> dict:
    keys = {}
    pl_path = out_dir / "translations_pl.json"
    if pl_path.is_file():
        try:
            data = json.loads(pl_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            data = None
        if isinstance(data, dict):
            for qid, question in data.items():
                if isinstance(question, dict):
                    keys[str(qid)] = polish_question_key(question)
    return keys


def merge_existing_translations(store, path: Path, unique_ids: set[str], changed_ids=None) -> None:
    if not path.is_file():
        return
    try:
        existing = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return
    if not isinstance(existing, dict):
        return
    changed_ids = changed_ids or set()
    for qid, tr in existing.items():
        qid = str(qid)
        if qid not in unique_ids or qid in changed_ids or qid in store:
            continue
        if not isinstance(tr, dict):
            continue
        kept = {k: v for k, v in tr.items() if k in ("q", "a", "b", "c") and v}
        if kept:
            src = tr.get("src")
            kept["src"] = src if src else "gemini"
            store[qid] = kept


def add_translation(store, qid, q, a, b, c, q_type):
    if not qid or qid in store:
        return
    if not q and not a and not b and not c:
        return
    tr = {}
    if q:
        tr["q"] = q
    if q_type == "specialist":
        if a:
            tr["a"] = a
        if b:
            tr["b"] = b
        if c:
            tr["c"] = c
    if tr:
        tr["src"] = "excel"
        store[qid] = tr


def load_pipeline(name: str):
    path = Path(__file__).with_name(name)
    spec = importlib.util.spec_from_file_location(name.replace(".", "_"), path)
    if spec is None or spec.loader is None:
        sys.exit(f"Cannot load {path}")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def load_download_gov():
    return load_pipeline("download-gov.py")


def main():
    dg = load_download_gov()
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--excel", default=str(dg.gov_data_dir() / "baza_pytan.xlsx"))
    parser.add_argument("--out-dir", default="src/data")
    parser.add_argument("--media-dir", default="")
    parser.add_argument(
        "--drop-missing-media",
        action="store_true",
        help="Clear media when the local raw file is missing. Default keeps Excel names for CDN.",
    )
    parser.add_argument(
        "--skip-translate-gaps",
        action="store_true",
        help="Do not fill missing EN/DE/UK after parsing. Default fills gaps.",
    )
    parser.add_argument(
        "--gemini-api-key",
        dest="gemini_api_key",
        help="Google AI Studio key for this run (overrides .geminienv). Omit to use .geminienv or Google Translate.",
    )
    args = parser.parse_args()

    project_root = Path(__file__).resolve().parent.parent
    excel_path = Path(args.excel) if os.path.isabs(args.excel) else project_root / args.excel
    out_dir = Path(args.out_dir) if os.path.isabs(args.out_dir) else project_root / args.out_dir
    media_dir = None
    if args.media_dir:
        media_dir = Path(args.media_dir) if os.path.isabs(args.media_dir) else project_root / args.media_dir
    elif args.drop_missing_media:
        preferred = dg.gov_raw_temp_dir()
        legacy = project_root / "Pytania egzaminacyjne na prawo jazdy 2025"
        media_dir = preferred if preferred.exists() else legacy

    dg.copy_legacy_app_gov_cache()
    if not excel_path.exists():
        sys.exit(
            f"Excel file not found: {excel_path}\n"
            "The question bank in git is already in src/data. To refresh that checkout from the ministry,\n"
            "download the Excel into the cache (no local server required):\n"
            "  python3 scripts/download-gov.py --excel-only\n"
            "then parse into this checkout:\n"
            "  python3 scripts/parse-excel.py\n"
            "--sync-gov parses into the local server data folder and does not write this checkout."
        )

    lookup = media_lookup(media_dir) if args.drop_missing_media else None
    out_dir.mkdir(parents=True, exist_ok=True)

    print(f"Loading {excel_path} ...")
    wb = openpyxl.load_workbook(str(excel_path), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if len(rows) < 2:
        sys.exit("Excel has no data rows.")

    header = list(rows[0])
    col_num = find_header(header, "Numer pytania", exact=["Numer pytania"])
    col_q = find_header(header, "Pytanie", exact=["Pytanie"])
    col_a = find_header(header, "Odpowiedz A", regex=r"^Odpowied.+ A$")
    col_b = find_header(header, "Odpowiedz B", regex=r"^Odpowied.+ B$")
    col_c = find_header(header, "Odpowiedz C", regex=r"^Odpowied.+ C$")
    col_correct = find_header(header, "Poprawna odp", exact=["Poprawna odp"])
    col_media = find_header(header, "Media", exact=["Media"])
    col_structure = find_header(header, "Zakres struktury", exact=["Zakres struktury"])
    col_points = find_header(header, "Liczba punktow", regex=r"^Liczba punkt", required=False)
    col_cats = find_header(header, "Kategorie", exact=["Kategorie"])
    col_en_q = find_header(header, "Pytanie [EN]", regex=r"^Pytanie \[EN\]$", required=False)
    col_en_a = find_header(header, "Odpowied A [EN]", regex=r"^Odpowied.+ A \[EN\]$", required=False)
    col_en_b = find_header(header, "Odpowied B [EN]", regex=r"^Odpowied.+ B \[EN\]$", required=False)
    col_en_c = find_header(header, "Odpowied C [EN]", regex=r"^Odpowied.+ C \[EN\]$", required=False)
    col_de_q = find_header(header, "Pytanie [D]", regex=r"^Pytanie \[D\]$", required=False)
    col_de_a = find_header(header, "Odpowied A [D]", regex=r"^Odpowied.+ A \[D\]$", required=False)
    col_de_b = find_header(header, "Odpowied B [D]", regex=r"^Odpowied.+ B \[D\]$", required=False)
    col_de_c = find_header(header, "Odpowied C [D]", regex=r"^Odpowied.+ C \[D\]$", required=False)
    col_uk_q = find_header(header, "Pytanie [UA]", regex=r"^Pytanie \[UA\]$", required=False)
    col_uk_a = find_header(header, "Odpowied A [UA]", regex=r"^Odpowied.+ A \[UA\]$", required=False)
    col_uk_b = find_header(header, "Odpowied B [UA]", regex=r"^Odpowied.+ B \[UA\]$", required=False)
    col_uk_c = find_header(header, "Odpowied C [UA]", regex=r"^Odpowied.+ C \[UA\]$", required=False)

    cat_questions = {cat: [] for cat in CATEGORIES}
    translations = {"en": {}, "de": {}, "uk": {}}
    unique_ids = set()
    missing_media = 0
    forced_specialist = 0

    for row in rows[1:]:
        row = list(row)
        qnum = cell(row, col_num)
        qtext = cell(row, col_q)
        correct = cell(row, col_correct)
        structure = cell(row, col_structure)
        raw_cats = cell(row, col_cats)
        raw_media = cell(row, col_media)
        ans_a = cell(row, col_a)
        ans_b = cell(row, col_b)
        ans_c = cell(row, col_c)
        raw_points = cell(row, col_points)

        q_type = question_type(structure, correct, ans_a, ans_b, ans_c)
        if q_type == "specialist" and structure == "PODSTAWOWY":
            forced_specialist += 1

        points = 0
        m = re.match(r"^\d+", raw_points)
        if m:
            points = int(m.group(0))
        if points < 1 or points > 3:
            points = 1 if q_type == "basic" else 2

        media_name, media_type, dropped = resolve_media(
            raw_media, lookup, args.drop_missing_media
        )
        if raw_media and dropped:
            missing_media += 1

        q_obj = {
            "id": int(qnum) if qnum.isdigit() else qnum,
            "q": qtext,
            "type": q_type,
            "correct": correct,
            "points": points,
            "media": media_name,
            "mediaType": media_type,
        }
        if q_type == "specialist":
            q_obj["a"] = ans_a
            q_obj["b"] = ans_b
            q_obj["c"] = ans_c

        qid = str(q_obj["id"]) if q_obj["id"] != "" else ""
        if qid:
            unique_ids.add(qid)
        add_translation(
            translations["en"], qid,
            cell(row, col_en_q), cell(row, col_en_a), cell(row, col_en_b), cell(row, col_en_c), q_type,
        )
        add_translation(
            translations["de"], qid,
            cell(row, col_de_q), cell(row, col_de_a), cell(row, col_de_b), cell(row, col_de_c), q_type,
        )
        add_translation(
            translations["uk"], qid,
            cell(row, col_uk_q), cell(row, col_uk_a), cell(row, col_uk_b), cell(row, col_uk_c), q_type,
        )

        for cat in raw_cats.split(","):
            cat = cat.strip()
            if cat in cat_questions:
                cat_questions[cat].append(q_obj)

    print(f"  {len(rows) - 1} Excel rows")
    if missing_media:
        print(f"  WARNING: {missing_media} questions lost media names (local file missing)")

    old_polish = load_existing_polish_keys(out_dir)
    pending = [q for cat in CATEGORIES for q in cat_questions[cat]]
    changed_ids = changed_polish_ids(old_polish, pending)
    if changed_ids:
        print(f"Polish text changed for {len(changed_ids)} questions; old translations of those ids will be filled again.")

    meta_categories = []
    total = 0
    question_facts = {}
    polish_texts = {}
    print()
    for cat in CATEGORIES:
        questions = cat_questions[cat]
        ids = []
        for question in questions:
            ids.append(question["id"])
            qid = str(question.get("id") or "")
            if not qid or qid in question_facts:
                continue
            question_facts[qid] = question_fact(question)
            polish_texts[qid] = polish_record(question)
        basic = sum(1 for q in questions if q["type"] == "basic")
        specialist = sum(1 for q in questions if q["type"] == "specialist")
        with open(out_dir / f"{cat}.json", "w", encoding="utf-8") as f:
            json.dump({"category": cat, "questionIds": ids}, f, ensure_ascii=False, indent=2)
            f.write("\n")
        meta_categories.append({
            "id": cat,
            "name": f"Kategoria {cat}",
            "questionCount": len(questions),
            "basicCount": basic,
            "specialistCount": specialist,
        })
        total += len(questions)
        print(f"  {cat:>3}: {len(questions):>4} questions ({basic} basic + {specialist} specialist)")

    meta = {
        "uniqueQuestionCount": len(unique_ids),
        "categories": meta_categories,
        "exam": EXAM_RULES,
    }
    with open(out_dir / "meta.json", "w", encoding="utf-8") as f:
        json.dump(meta, f, ensure_ascii=False, indent=2)
        f.write("\n")
    print(f"Wrote meta.json ({len(unique_ids)} unique, {total} category assignments).")

    for filename, payload in (
        ("questions.json", question_facts),
        ("translations_pl.json", polish_texts),
    ):
        with open(out_dir / filename, "w", encoding="utf-8") as f:
            json.dump(payload, f, ensure_ascii=False, separators=(",", ":"))
        print(f"Wrote {filename} ({len(payload)} questions).")

    for lang, filename in (("en", "translations_en.json"), ("de", "translations_de.json"), ("uk", "translations_uk.json")):
        merge_existing_translations(translations[lang], out_dir / filename, unique_ids, changed_ids)
        with open(out_dir / filename, "w", encoding="utf-8") as f:
            json.dump(translations[lang], f, ensure_ascii=False, separators=(",", ":"))
        print(f"Wrote {filename} ({len(translations[lang])} questions).")

    if not args.skip_translate_gaps:
        print("Filling translation gaps (Gemini if a key is available, else Google Translate; --skip-translate-gaps to skip)...")
        load_pipeline("translate-questions.py").fill_missing(
            out_dir, gemini_api_key=args.gemini_api_key
        )

    if forced_specialist:
        print(f"Treated {forced_specialist} PODSTAWOWY rows as specialist (correct answer A/B/C).")
    print("Done.")


if __name__ == "__main__":
    main()
